"""VMS · Welcome Pack web · v4 — TODO EL CONTENIDO SALE DEL EXCEL  08_WELCOME_WEB/_datos/VMS_contenido.xlsx
Cambios v4 (pedidos por Miguel 2026-10-01): portada con horarios + Cómo llegar / Llamar / WhatsApp · menú visual «Tu guía» ·
2 redes Wi-Fi con QR · Maps directo en «A mano» · galería por zona con carrusel en «Tenemos para ti» · «No olvides traer» visible ·
iconos en todas las reglas · fichas de actividad/lugar con foto y descripción · directorio con filtros · panel de pendientes.
Uso: python3 99_SCRIPTS/web_build_v4.py  (o doble clic en 08_WELCOME_WEB/ACTUALIZAR_WEB.command)  ->  08_WELCOME_WEB/sitio_v4/vms/
Revisión: vms/?pendientes#inicio  (avisos amarillos + botón «Pendientes»)."""
import os, sys, re, json, shutil, html, urllib.request
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import web_build_v2 as v2
import web_build_v3 as v3
from openpyxl import load_workbook
e = html.escape
R, W = v2.R, v2.W
OUT = os.path.join(W, "sitio_v4"); FOT = os.path.join(W, "_fuente", "fotos")
# Contenido: 1) variable VMS_XLSX  2) el Excel de Google Drive sincronizado en este Mac  3) la copia de _datos (la que usa GitHub)
XLSX = next((p for p in (os.environ.get("VMS_XLSX", ""),
             os.path.expanduser("~/Library/CloudStorage/GoogleDrive-miguel@nadena.net/My Drive/VMS_WEB_CONTENIDO/VMS_contenido.xlsx"))
             if p and os.path.exists(p)), os.path.join(W, "_datos", "VMS_contenido.xlsx"))
VENDOR = os.path.join(W, "_fuente", "vendor")
I = v2.I
I.update({
 "shower": '<path d="M4 21V9a5 5 0 0 1 10 0"/><path d="M10 9h8l-1.5 3h-5z"/><path d="M12 15v1M15 15v1M18 15v1M13.5 18v1M16.5 18v1"/>',
 "lifebuoy": '<circle cx="12" cy="12" r="9"/><circle cx="12" cy="12" r="4"/><path d="M5.6 5.6l3.6 3.6M14.8 14.8l3.6 3.6M18.4 5.6l-3.6 3.6M9.2 14.8l-3.6 3.6"/>',
 "ban": '<circle cx="12" cy="12" r="9"/><path d="M5.6 5.6l12.8 12.8"/>',
 "thermo": '<path d="M14 14.8V5a2 2 0 0 0-4 0v9.8a4 4 0 1 0 4 0z"/><path d="M12 9v7"/>',
 "wa": '<path d="M4 20l1.4-4.2A8 8 0 1 1 8.8 19z"/><path d="M9.2 9.3c.2 2.6 2.5 4.9 5.2 5.2l1.1-1.3-1.8-.9-.9.8a3.6 3.6 0 0 1-1.9-1.9l.8-.9-.9-1.8z"/>',
 "nav": '<path d="M3 11l18-8-8 18-2-8z"/>',
 "clock": '<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 2"/>',
 "bag2": '<path d="M6 7h12l1 14H5zM9 7a3 3 0 0 1 6 0"/>',
})
def ic(k, cls="i"): return v2.ic(k if k in I else "info", cls)
fs = v2.fs
def yes(v): return str(v).strip().lower() in ("sí", "si", "yes", "x", "1", "true", "s")

# ---------------- carga del Excel ----------------
D = {}; G = {}; T = {}; PEND = []; ASSETS_COPIED = set()
def num(v):
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return "" if v is None else str(v).strip()
def load():
    if not os.path.exists(XLSX): sys.exit("No existe el Excel. Ejecuta antes: python3 99_SCRIPTS/crear_excel_contenido.py")
    wb = load_workbook(XLSX, data_only=True)
    for n in wb.sheetnames:
        if n in ("LEEME", "Iconos"): continue
        ws = wb[n]; hdr = [num(c.value) for c in ws[1]]; rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            vals = [num(x) for x in r][:len(hdr)]
            if any(vals): rows.append(dict(zip(hdr, vals + [""] * (len(hdr) - len(vals)))))
        D[n] = rows
        for i, row in enumerate(rows):
            for k, v in row.items():
                if "PENDIENTE" in v.upper() or "BORRADOR" in v.upper():
                    ref = row.get("clave") or row.get("nombre") or row.get("red") or row.get("categoria") or row.get("etiqueta") or f"fila {i + 2}"
                    PEND.append(f"<b>{e(n)} · {e(ref)}</b>: {e(v)}")
    G.update({r["clave"]: r["valor"] for r in D["General"]}); T.update({r["clave"]: r["texto"] for r in D["Textos"]})
def miss(k): PEND.append(f"<b>Falta</b> el dato «{e(k)}»"); return f'<span class="miss">[{e(k)}]</span>'
def g(k): return e(G[k]) if G.get(k) else miss(k)
def t(k): return e(T[k]) if T.get(k) else miss(k)
def graw(k): return G.get(k, "")

# ---------------- fotos ----------------
def urls(field, who):
    out = []
    for p in [x.strip() for x in (field or "").replace(",", ";").split(";") if x.strip()]:
        src = os.path.join(FOT, p)
        if os.path.exists(src):
            dst = os.path.join(OUT, "assets", "fotos", p); os.makedirs(os.path.dirname(dst), exist_ok=True)
            if p not in ASSETS_COPIED: shutil.copy(src, dst); ASSETS_COPIED.add(p)
            out.append("../assets/fotos/" + p)
        else: PEND.append(f"<b>Foto no encontrada</b> «{e(p)}» ({e(who)})")
    return out
def car(us, label, h=""):
    sl = "".join(f'<div class="sl" style="background-image:url({u})"></div>' for u in us) or f'<div class="sl ph" data-ph="Foto pendiente · {e(label)}"></div>'
    cnt = f'<span class="cnt">1 / {len(us)}</span>' if len(us) > 1 else ""
    return f'<div class="car"{f" style={chr(34)}--h:{h}{chr(34)}" if h else ""}><div class="slides">{sl}</div>{cnt}</div>'
def credf(s):
    return f'<p class="credf">Foto: {e(s)}</p>' if s else ""
def foto(slot, cls, inner, desc, extra=""):
    p = os.path.join(FOT, slot + ".jpg")
    if os.path.exists(p):
        shutil.copy(p, os.path.join(OUT, "assets", "fotos", slot + ".jpg"))
        return f'<div class="{cls}" {extra} style="background-image:url(../assets/fotos/{slot}.jpg)"><div class="vel"></div>{inner}</div>'
    PEND.append(f"<b>Foto pendiente</b>: {e(desc)} (_fuente/fotos/{slot}.jpg)")
    return f'<div class="{cls} ph" {extra} data-ph="Foto pendiente · {e(desc)}"><div class="vel"></div>{inner}</div>'

# ---------------- piezas ----------------
CHEV = '<span class="cv">' + ic("chev") + "</span>"
def eyebrow(x): return f'<div class="eb"><i></i><span>{x}</span></div>'
def header(fam, sec, title, intro): return f'<header class="hd" {fs(fam)}>{eyebrow(sec)}<h1>{title}</h1><p>{intro}</p></header>'
def stitle(x, det="", sheet="", fam=None, id_=""):
    d = (f'<a class="det" href="#" data-sheet="{sheet}">{det}</a>' if sheet else f'<span class="det">{det}</span>') if det else ""
    return f'<div class="st"{f" id={chr(34)}{id_}{chr(34)}" if id_ else ""}{" " + fs(fam) if fam else ""}><h2>{x}</h2>{d}</div>'
def chip(icon, size=""): return f'<span class="chip {size}">{ic(icon)}</span>'
def punto(icon, text): return f'<li class="pt">{chip(icon)}<span>{text}</span></li>'
def rev(x): return f'<p class="rev"><b>Pendiente:</b> {x}</p>'
def tel_href(n):
    d = "".join(c for c in n if c.isdigit())
    return "tel:+" + (d if len(d) == 11 else "1" + d) if d else ""
def wa_href(): return f'https://wa.me/{"".join(c for c in graw("whatsapp") if c.isdigit())}?text={urllib.request.quote(graw("whatsapp_mensaje"))}'
PILL = [("telefono", "tel", "Llamar"), ("maps", "nav", "Cómo llegar"), ("instagram", "ig", "Instagram"), ("web", "web", "Web"), ("tripadvisor", "trip", "TripAdvisor"), ("wikiloc", "ruta", "Rutas en Wikiloc")]
def pills(p):
    out = ""
    for k, icon, lab in PILL:
        v = p.get(k, "")
        if not v: continue
        href = tel_href(v) if k == "telefono" else v
        out += f'<a class="pill" href="{e(href)}"{"" if k == "telefono" else " target=_blank rel=noopener"}>{ic(icon)}{lab}</a>'
    return f'<div class="pills">{out}</div>' if out else ""
def quick(p): return v2.acts(p["nombre"], tel=p.get("telefono", ""), maps=p.get("maps", ""), ig=p.get("instagram", ""), web=p.get("web", ""), trip=p.get("tripadvisor", ""), ruta=p.get("wikiloc", ""))
def slug(s): return re.sub(r"[^a-z0-9]+", "-", s.lower().translate(str.maketrans("áéíóúñü", "aeiouñu")).replace("ñ", "n")).strip("-")

# ---------------- PORTADA ----------------
def s_portada():
    inner = (f'<div class="marca"><span class="sello"><i class="mark"></i></span><span class="tipo">Welcome book</span></div>'
             f'<div class="pres"><p class="ubi">{ic("pin")}{g("ubicacion")}</p>'
             f'<div class="tit"><h1>{g("nombre")}</h1><p class="prom">{g("promesa")}</p></div>'
             f'<div class="datos"><div><b>{g("camas")}</b><span>camas</span></div><div><b>{g("banos")}</b><span>baños</span></div><div><b>{g("huespedes")}</b><span>huéspedes</span></div></div>'
             f'<a class="horas" href="#casa:checkin"><div><span>{ic("in")}Check-in</span><b>{g("checkin")}</b></div><div><span>{ic("out")}Check-out</span><b>{g("checkout")}</b></div>{ic("chev")}</a>'
             f'<div class="rap"><a href="{e(graw("maps"))}" target="_blank" rel="noopener">{ic("nav")}Cómo llegar</a>'
             f'<a class="urg" href="{tel_href(graw("telefono_llamar"))}">{ic("tel")}Llamar</a>'
             f'<a href="{e(wa_href())}" target="_blank" rel="noopener">{ic("wa")}WhatsApp</a></div>'
             f'<a class="abrir" href="#inicio"><span>Abrir guía</span><i>{ic("arrow")}</i></a></div>')
    return f'<section data-view="portada" class="portada" {fs("bienvenida")}>' + foto("portada", "pt-bg", inner, "exterior de la villa") + "</section>"

# ---------------- INICIO ----------------
FAMV = {"inicio": "bienvenida", "casa": "llegada", "reglas": "casa", "explora": "jarabacoa"}
def guia():
    out = ""
    for r in D["Guia"]:
        dest = r["destino"]; v = re.split("[/:]", dest)[0]
        fam = "emergencias" if "emergencias" in dest else ("wifi" if "wifi" in dest else FAMV.get(v, "bienvenida"))
        out += f'<a class="gi" href="#{e(dest)}" {fs(fam)}><span class="ico">{ic(r["icono"])}</span>{e(r["etiqueta"])}</a>'
    return f'<nav class="guia" aria-label="{t("guia_titulo")}">{out}</nav>'
def wifi_cards():
    out = ""
    for w in D["WiFi"]:
        s = lambda x: re.sub(r'([\\;,:"])', r"\\\1", x)
        payload = f'WIFI:T:{w.get("seguridad") or "WPA"};S:{s(w["red"])};P:{s(w["contraseña"])};;'
        out += (f'<article class="card wifi" {fs("wifi")}><div class="cab">{chip("wifi", "l")}<span class="estado">{e(w.get("etiqueta", ""))}</span></div>'
                f'<div class="row2"><div class="creds"><button class="cp" data-copy="{e(w["red"])}"><span class="lb">Red</span><b>{e(w["red"])}</b></button>'
                f'<button class="cp" data-copy="{e(w["contraseña"])}"><span class="lb">Contraseña</span><b>{e(w["contraseña"])}</b></button></div>'
                f'<div class="qr" data-qr="{e(payload)}" role="img" aria-label="Código QR de la red {e(w["red"])}"></div></div></article>')
    return out
def s_inicio():
    emer_btn = f'<div class="btns sm"><a class="btn urg" href="{tel_href(graw("emergencias_telefono"))}">{ic("tel")}Llamar a emergencias</a></div>' if graw("emergencias_telefono") else ""
    return f'''<section data-view="inicio" class="view">
{header("bienvenida", t("inicio_seccion"), t("inicio_titulo"), t("inicio_intro"))}
{stitle(t("guia_titulo"), "Desliza →")}{guia()}
<article class="card host tap" data-sheet="anfitriones" {fs("bienvenida")}>{foto("anfitriones", "retrato", "", "Julián y Deysi")}
 <div class="msg"><div class="firma"><h3>{g("anfitriones")}</h3><span class="badge">Tus anfitriones</span></div><p>{t("bienvenida")} <span class="more">Leer más</span></p></div></article>
<div class="sec" id="wifi">{stitle("Conéctate", t("wifi_estado"))}{wifi_cards()}</div>
<div class="sec">{stitle("A mano")}
 <article class="card info tap" data-sheet="contacto" {fs("bienvenida")}>{chip("tel", "l")}<div><span class="lb">Contacto</span><b>{g("anfitriones").split(" y ")[0]} · {g("telefono")}</b><p>{g("email")} · Propietario VMS</p></div>{CHEV}</article>
 <article class="card info tap" data-sheet="llegar" {fs("bienvenida")}>{chip("pin", "l")}<div><span class="lb">Ubicación</span><b>{g("direccion")}</b><p>{g("zona")}</p>
  <div class="btns sm"><a class="btn solid" href="{e(graw("maps"))}" target="_blank" rel="noopener">{ic("nav")}Abrir en Google Maps</a></div></div>{CHEV}</article>
 <article class="card info tint tap" data-sheet="emergencias" {fs("emergencias")}>{chip("siren", "l")}<div><span class="lb">Emergencias</span>
  <b class="{"" if graw("emergencias_telefono") else "falta"}">{e(graw("emergencias_telefono")) or "Teléfono pendiente"}</b><p>{t("emergencias_detalle")}</p>{emer_btn}</div>{CHEV}</article></div>
<article class="card seg tint tap" data-sheet="emergencias" {fs("emergencias")}><div class="cab2">{ic("shield")}<h4>{t("seguridad_titulo")}</h4>{CHEV}</div>
 <ul class="pts">{punto("shield", t("botiquin"))}{punto("flame", t("extintores"))}</ul></article>
</section>'''

# ---------------- LA CASA ----------------
AMEN = []
def s_casa():
    def pasos(tipo): return [r["texto"] for r in D["Pasos"] if r["tipo"] == tipo and yes(r.get("en_tarjeta", "sí"))]
    def horario(fam, icon, lab, hora, ps, id_):
        return (f'<article class="card hor tap" id="{id_}" data-sheet="llegada" {fs(fam)}><div class="cab"><span class="tipo">{ic(icon)}{lab}</span><b class="hora">{hora}</b></div>'
                f'<hr><ul class="pts">' + "".join(punto("check", e(p)) for p in ps) + "</ul></article>")
    basic = [r["item"] for r in D["Traer"] if r.get("tipo", "básico") != "actitud"]; act = [r["item"] for r in D["Traer"] if r.get("tipo") == "actitud"]
    traer = (f'<article class="card traer" id="traer" {fs("debes")}><div class="tema">{chip("bag2", "m")}<div><h3>{t("traer_titulo")}</h3><p class="sub">{t("traer_intro")}</p></div></div>'
             f'<div class="tags">' + "".join(f"<span>{e(x)}</span>" for x in basic) + '</div><div class="tags plus">' + "".join(f"<span>{e(x)}</span>" for x in act) + "</div></article>")
    am = ""
    for i, r in enumerate(D["Amenidades"]):
        AMEN.append({"l": r["etiqueta"], "i": ic(r["icono"]), "d": r.get("descripcion", ""), "f": urls(r.get("fotos", ""), r["etiqueta"])})
        if yes(r.get("en_resumen", "sí")):
            am += f'<a class="am" href="#" data-sheet="galeria.{i}">{ic(r["icono"])}<span>{e(r["etiqueta"])}</span></a>'
    if not any(a["f"] for a in AMEN): PEND.append("<b>Galería «Tenemos para ti»</b>: no hay fotos de ninguna zona (columna fotos de la hoja Amenidades).")
    exp = foto("casa", "photo exp", f'<span class="badge glass">{t("foto_casa_distintivo")}</span><p class="frase">{t("foto_casa_frase")}</p>', "piscina o terraza")
    return f'''<section data-view="casa" class="view">
{header("llegada", t("casa_seccion"), t("casa_titulo"), t("casa_intro"))}
<div class="sec">{stitle("Check-in / out", "Ver detalles", "llegada", "llegada")}
 {horario("llegada", "in", "Check-in", g("checkin"), pasos("checkin"), "checkin")}
 {horario("casa", "out", "Check-out", g("checkout"), pasos("checkout"), "checkout")}
 {traer}</div>
<article class="card feat info tap" data-sheet="debes" {fs("debes")}>{chip("users", "l")}<div><span class="lb">Personal de apoyo</span><b>{g("personal_horario")}</b><p>{t("personal_texto")}</p></div>{CHEV}</article>
{exp}
<div class="sec" {fs("casa")}>{stitle(t("amenidades_titulo"), "Inventario completo", "inventario", "casa", "tenemos")}
 <p class="hint">{ic("sparkles")}Toca una zona para ver sus fotos.</p>
 <div class="amen">{am}</div><p class="nota">{t("amenidades_nota")}</p></div>
</section>'''

# ---------------- REGLAS ----------------
def s_reglas():
    def grupo(fam, icon, title, key, dato="", sheet=""):
        items = [(r["icono"], e(r["texto"])) for r in D["Reglas_resumen"] if r["tarjeta"] == key]
        d = f'<span class="dato">{dato}</span>' if dato else ""
        return (f'<article class="card grp tap" data-sheet="{sheet}" {fs(fam)}><div class="cab"><div class="tema">{chip(icon, "m")}<h3>{title}</h3></div>{d}{CHEV}</div><hr>'
                '<ul class="pts">' + "".join(punto(i, x) for i, x in items) + "</ul></article>")
    agua = foto("piscina", "photo agua", f'<span class="lbl">{t("agua_etiqueta")}</span><p class="frase">{t("agua_titulo")}</p>', "piscina o jacuzzi")
    return f'''<section data-view="reglas" class="view">
{header("casa", t("reglas_seccion"), t("reglas_titulo"), t("reglas_intro"))}
<article class="card feat cap tap" data-sheet="reglas-casa" {fs("casa")}><div class="cifra"><b>{g("huespedes")}</b><span>máximo</span></div><div class="cond"><b>{t("capacidad_titulo")}</b><p>{t("capacidad_texto")}</p></div>{CHEV}</article>
{grupo("casa", "housepl", "En la casa", "casa", sheet="reglas-casa")}
{agua}
{grupo("piscina", "waves", "Piscina", "piscina", t("dato_piscina"), "reglas-piscina")}
{grupo("piscina", "bath", "Jacuzzi", "jacuzzi", t("dato_jacuzzi"), "reglas-piscina")}
{grupo("casa", "pawprint", "Mascotas", "mascotas", t("dato_mascotas"), "reglas-mascotas")}
</section>'''

# ---------------- EXPLORA ----------------
def lugares(sec=None, cat=None, res=None):
    return [(i, p) for i, p in enumerate(D["Lugares"]) if (sec is None or p["seccion"] == sec) and (cat is None or p["categoria"] == cat) and (res is None or yes(p.get("en_resumen")) == res)]
def cats(sec): return [c for c in D["Categorias"] if c["seccion"] == sec]
def s_explora():
    dest = foto("explora", "photo dest tap", f'<span class="badge glass">{ic("leaf")}{t("destacado_etiqueta")}</span><div class="txt"><p class="frase">{t("destacado_titulo")}</p><p>{t("destacado_texto")}</p>{(chr(60) + "small class=credd>Foto: " + e(T.get("destacado_credito", "")) + chr(60) + "/small>") if T.get("destacado_credito") else ""}</div>', "Jarabacoa", 'data-sheet="sec-naturaleza"')
    def card(icon, title, pills_, sheet, tint=False):
        return (f'<article class="card cat{" tint" if tint else ""}" {fs("jarabacoa")}><a class="cab" href="#" data-sheet="{sheet}">{chip(icon, "m")}<h3>{title}</h3><span class="todo">Ver todo</span>{CHEV}</a>'
                '<div class="ops">' + "".join(f'<a class="lugar" href="#" data-sheet="{s}">{ic("pin")}{e(n)}</a>' for n, s in pills_) + "</div></article>")
    nat = [(p["nombre"], f"l{i}") for i, p in lugares("naturaleza", res=True)]
    av = [(c["categoria"], "c-" + slug(c["categoria"])) for c in cats("aventura")]
    tq = [(p["nombre"], f"l{i}") for i, p in lugares("tranquilos", res=True)]
    rest = "".join(f'<a class="rr" href="#" data-sheet="l{i}"><span>{e(p["nombre"])}</span>{ic("chev")}</a>' for i, p in lugares("comer", res=True))
    dirs = "".join(f'<a class="dr" href="#" data-sheet="l{i}"><b>{e(p["nombre"])}</b><span>{e(p["categoria"])}{" · delivery" if "delivery" in p.get("subtitulo", "").lower() else ""}</span></a>' for i, p in lugares("directorio", res=True))
    n_comer = len(lugares("comer"))
    return f'''<section data-view="explora" class="view">
{header("jarabacoa", t("explora_seccion"), t("explora_titulo"), t("explora_intro"))}
{dest}
<div class="sec">{stitle("Experiencias", "Para todos")}
 {card("trees", "Naturaleza", nat, "sec-naturaleza")}
 {card("mountain", "Aventura", av, "sec-aventura", True)}
 {card("sun", "Planes tranquilos", tq, "sec-tranquilos")}</div>
<div class="sec" id="comer" {fs("jarabacoa")}>{stitle(t("comer_titulo"), f"Ver los {n_comer}", "sec-comer", "jarabacoa")}<div class="card rest" {fs("jarabacoa")}>{rest}</div></div>
<div class="sec" id="directorio">{stitle(t("directorio_titulo"), "Ver todo", "directorio", "jarabacoa")}<div class="card dir" {fs("jarabacoa")}>{dirs}</div></div>
</section>'''

# ---------------- HOJAS ----------------
def lst(items): return "<ul class='ul'>" + "".join(f"<li>{x}</li>" for x in items) + "</ul>"
def blk(x, body): return f'<div class="blk"><h4>{x}</h4>{body}</div>'
def btn(href, icon, label, solid=False, ext=True, sheet=""):
    a = f' data-sheet="{sheet}"' if sheet else (" target=_blank rel=noopener" if ext else "")
    return f'<a class="btn{" solid" if solid else ""}" href="{e(href)}"{a}>{ic(icon)}{label}</a>'
def row_link(i, p, show_cat=False):
    sub = " · ".join(x for x in ((p["categoria"] if show_cat else ""), p.get("subtitulo", "")) if x)
    return (f'<div class="pl{" star" if yes(p.get("destacado")) else ""}" data-cat="{e(p["categoria"])}"><a class="pn" href="#" data-sheet="l{i}"><b>{e(p["nombre"])}</b>'
            f'{f"<small>{e(sub)}</small>" if sub else ""}</a>{quick(p)}</div>')
SEC_T = {"naturaleza": ("Qué hacer", "Naturaleza"), "aventura": ("Qué hacer", "Aventura"), "tranquilos": ("Qué hacer", "Planes tranquilos"), "comer": ("Jarabacoa", "Dónde comer y beber")}
def sheets():
    S = {}
    rc = graw("instagram")
    S["anfitriones"] = ("bienvenida", "Tus anfitriones", g("anfitriones"),
        f"<p>{t('bienvenida')}</p><p>{t('bienvenida_2')}</p>" + blk("Acerca de Villa Mi Sueño", f"<p>{t('acerca_1')}</p><p>{t('acerca_2')}</p><p>{t('acerca_3')}</p>")
        + blk("Comparte tu estancia", f"<p>{t('comparte_instagram')}</p><div class='btns'>{btn('https://www.instagram.com/' + rc + '/', 'ig', '@' + e(rc))}</div><p>{t('comparte_review')}</p>")
        + blk(t("club_titulo"), f"<div class='dots'><i class='ok'></i><i class='ok'></i><i class='ok'></i><i class='ok'></i><i></i></div><p>{t('club_texto')}</p><div class='btns'>"
              + btn("mailto:" + graw("email") + "?subject=Reserva%20Villa%20Mi%20Sue%C3%B1o", "mail", "Escríbenos: " + g("email"), True, False) + "</div>"))
    S["contacto"] = ("bienvenida", "Contacto", g("propietario"),
        "<p>Propietario de Villa Mi Sueño. Escríbenos o llámanos para lo que necesites.</p><div class='btns col'>"
        + btn(tel_href(graw("telefono_llamar")), "tel", "Llamar · " + g("telefono"), True, False) + btn(wa_href(), "wa", "WhatsApp") + btn("mailto:" + graw("email"), "mail", g("email"), ext=False) + "</div>")
    S["llegar"] = ("llegada", "Cómo llegar", "Ruta a Villa Mi Sueño",
        f"<p><b>{g('direccion')}, {g('zona')}, R.D.</b></p><div class='btns'>{btn(graw('maps'), 'nav', 'Abrir en Google Maps', True)}</div>"
        + blk("Si no carga el GPS", "<p>No importa: sigue las instrucciones al pie de la letra.</p><ol class='ruta'>" + "".join(f"<li>{e(r['paso'])}</li>" for r in D["Ruta"]) + "</ol>"))
    em = "".join(f'<div class="pl"><div class="pn"><b>{e(r["nombre"])}</b>{"" if (r.get("telefono") or r.get("maps")) else "<small class=falta>Teléfono y mapa pendientes</small>"}</div>'
                 f'{v2.acts(r["nombre"], tel=r.get("telefono", ""), maps=r.get("maps", ""))}</div>' for r in D["Emergencias"])
    S["emergencias"] = ("emergencias", "En caso de", "Emergencias", (f"<div class='btns'>{btn(tel_href(graw('emergencias_telefono')), 'tel', 'Llamar a emergencias · ' + g('emergencias_telefono'), True, False)}</div>" if graw("emergencias_telefono") else "") + blk("En la casa", lst([t("extintores"), t("botiquin")])) + blk("Servicios", f'<div class="card list">{em}</div>'))
    S["llegada"] = ("llegada", "Check-in / out", "Llegada y salida",
        blk(f"Check-in · {g('checkin')}", lst([e(r["texto"]) for r in D["Pasos"] if r["tipo"] == "checkin"] + [t("checkin_flexible")]))
        + blk(f"Check-out · {g('checkout')}", lst([e(r["texto"]) for r in D["Pasos"] if r["tipo"] == "checkout"] + [t("checkout_flexible")]))
        + f"<div class='btns'>{btn('#', 'nav', 'Cómo llegar · ruta paso a paso', sheet='llegar')}</div>")
    S["debes"] = ("debes", "Debes saber", "Para una estancia tranquila", "".join(blk(e(r["titulo"]), f"<p>{e(r['texto'])}</p>") for r in D["Debes"]))
    grp = {}
    for r in D["Inventario"]: grp.setdefault(r["grupo"], []).append(e(r["item"]))
    S["inventario"] = ("casa", "Tenemos para ti", "Todo lo que hay en la villa", "".join(blk(e(k), lst(v)) for k, v in grp.items()))
    for hoja, (fam, eb, tit, intro) in {"casa": ("casa", "Reglas", "Reglas de la casa", ""), "piscina": ("piscina", "Reglas", "Piscina y jacuzzi", ""),
                                        "mascotas": ("casa", "Somos pet friendly", "Mascotas", t("mascotas_intro"))}.items():
        gg = {}
        for r in D["Reglas_completas"]:
            if r["hoja"] == hoja: gg.setdefault(r["grupo"], []).append(e(r["texto"]))
        S["reglas-" + hoja] = (fam, eb, tit, (f"<p>{intro}</p>" if intro else "") + "".join(blk(e(k), lst(v)) for k, v in gg.items()))
    # galería (contenido dinámico por JS)
    zn = "".join(f'<button class="zn" data-z="{i}"><span>{a["i"]}</span>{e(a["l"])}</button>' for i, a in enumerate(AMEN))
    S["galeria"] = ("casa", "Tenemos para ti", "", f'<div class="gal">{car([], "zona")}<p class="gdesc"></p></div><div class="zonas-w"><p class="lb zt">Otras zonas</p><div class="zonas">{zn}</div></div>')
    # secciones de Explora
    for sec, (eb, tit) in SEC_T.items():
        S["sec-" + sec] = ("jarabacoa", eb, tit, '<div class="card list">' + "".join(row_link(i, p, sec == "aventura") for i, p in lugares(sec)) + "</div>")
    # categorías de aventura (p.ej. Parapente -> sus 3 empresas)
    for c in cats("aventura"):
        def op(i, p):
            sub = f"<small>{e(p['subtitulo'])}</small>" if p.get("subtitulo") else ""
            ds = f"<p>{e(p['descripcion'])}</p>" if p.get("descripcion") else ""
            return f'<article class="op"><a class="oph" href="#" data-sheet="l{i}"><h4>{e(p["nombre"])}</h4>{CHEV}</a>{sub}{ds}{pills(p)}</article>'
        ops = "".join(op(i, p) for i, p in lugares("aventura", c["categoria"]))
        n = len(lugares("aventura", c["categoria"]))
        trip = f"<div class='btns'>{btn(c['tripadvisor'], 'trip', 'Ver opiniones en TripAdvisor')}</div>" if c.get("tripadvisor") else ""
        S["c-" + slug(c["categoria"])] = ("jarabacoa", "Aventura", e(c["categoria"]),
            car(urls(c.get("fotos", ""), c["categoria"]), c["categoria"]) + credf(c.get("credito_fotos", "")) + (f"<p>{e(c['descripcion'])}</p>" if c.get("descripcion") else "")
            + (rev("descripción en borrador: verificar con los operadores.") if "BORRADOR" in c.get("nota", "").upper() else "")
            + f'<p class="lb cnt2">{n} {"opción" if n == 1 else "opciones"}</p>{ops}{trip}')
    # fichas de cada lugar
    for i, p in enumerate(D["Lugares"]):
        sec_l = {"naturaleza": "Naturaleza", "aventura": "Aventura · " + p["categoria"], "tranquilos": "Planes tranquilos", "comer": "Comer y beber", "directorio": p["categoria"]}.get(p["seccion"], p["categoria"])
        desc = f"<p>{e(p['descripcion'])}</p>" if p.get("descripcion") else rev("falta una descripción breve (columna «descripcion» de la hoja Lugares).")
        back = {"aventura": "c-" + slug(p["categoria"]), "directorio": "directorio"}.get(p["seccion"], "sec-" + p["seccion"])
        S[f"l{i}"] = ("jarabacoa", e(sec_l), e(p["nombre"]),
            car(urls(p.get("fotos", ""), p["nombre"]), p["nombre"]) + credf(p.get("credito_fotos", "")) + (f'<p class="subt">{e(p["subtitulo"])}</p>' if p.get("subtitulo") else "") + desc + pills(p)
            + f"<div class='btns'>{btn('#', 'chev', 'Ver más opciones', sheet=back)}</div>")
    # directorio con filtros
    dcats = []
    for i, p in lugares("directorio"):
        if p["categoria"] not in dcats: dcats.append(p["categoria"])
    filt = '<div class="filt"><button class="on" data-f="*">Todos</button>' + "".join(f'<button data-f="{e(c)}">{e(c)}</button>' for c in dcats) + "</div>"
    S["directorio"] = ("jarabacoa", "Jarabacoa", "Directorio local", filt + '<div class="card list">' + "".join(row_link(i, p, True) for i, p in lugares("directorio")) + "</div>")
    # pendientes
    S["pendientes"] = ("emergencias", "Modo revisión", f"Pendientes ({len(set(PEND))})", lst(sorted(set(PEND))))
    out = ""
    for k, (fam, eb, title, body) in S.items():
        cls = " gsheet" if k == "galeria" else ""
        out += (f'<div class="sheet{cls}" id="sh-{k}" role="dialog" aria-modal="true" hidden {fs(fam)}><div class="panel">'
                f'<div class="sh-h"><div>{eyebrow(eb)}<h2>{title}</h2></div><button class="x" aria-label="Cerrar">{ic("x")}</button></div>'
                f'<div class="sh-b">{body}</div></div></div>')
    return out

NAV = [("inicio", "Inicio", "home"), ("casa", "La casa", "key"), ("reglas", "Reglas", "clip"), ("explora", "Explora", "map")]
CSS4 = r"""
.pt-bg{padding-top:18px}.tit h1{font-size:42px;margin-top:14px}.prom{margin-bottom:16px}.datos{margin-bottom:10px;padding:12px 0}.datos b{font-size:23px}
.horas{display:flex;align-items:center;text-decoration:none;color:#fff;border:1px solid rgba(255,255,255,.2);background:rgba(255,255,255,.08);border-radius:16px;margin-bottom:10px;-webkit-backdrop-filter:blur(8px);backdrop-filter:blur(8px);padding-right:10px}
.horas div{flex:1;padding:10px 14px;display:flex;flex-direction:column;gap:5px}.horas div+div{border-left:1px solid rgba(255,255,255,.2)}
.horas span{display:flex;align-items:center;gap:6px;font:600 10px var(--txt);letter-spacing:1.4px;text-transform:uppercase;opacity:.85}.horas span .i{width:13px;height:13px}
.horas b{font:700 18px var(--disp)}.horas>.i{width:16px;height:16px;opacity:.6}
.rap{display:grid;grid-template-columns:repeat(3,1fr);gap:8px;margin-bottom:12px}
.rap a{display:flex;flex-direction:column;align-items:center;gap:6px;padding:11px 4px;border-radius:16px;background:rgba(255,255,255,.1);border:1px solid rgba(255,255,255,.18);color:#fff;text-decoration:none;font:600 11.5px var(--txt);-webkit-backdrop-filter:blur(8px);backdrop-filter:blur(8px)}
.rap a .i{width:20px;height:20px}.rap a.urg{background:rgba(158,27,20,.62);border-color:rgba(255,255,255,.22)}
.guia{display:flex;gap:10px;overflow-x:auto;scroll-snap-type:x proximity;margin:0 -20px 28px;padding:2px 20px 4px;scrollbar-width:none;-webkit-mask-image:linear-gradient(90deg,#000 calc(100% - 34px),transparent);mask-image:linear-gradient(90deg,#000 calc(100% - 34px),transparent)}
.guia::-webkit-scrollbar{display:none}
.gi{flex:none;width:66px;scroll-snap-align:start;display:flex;flex-direction:column;align-items:center;gap:7px;text-decoration:none;text-align:center;font:600 10.5px/1.2 var(--txt);color:var(--tinta)}
.gi .ico{width:50px;height:50px;border-radius:16px;display:grid;place-items:center;background:var(--chip);color:var(--ink);transition:transform .12s}.gi:active .ico{transform:scale(.94)}.gi .i{width:21px;height:21px}
.st .det{text-align:right}
.wifi .row2{display:flex;gap:16px;align-items:center;margin-top:16px}.creds{flex:1;display:grid;gap:14px;min-width:0}
.qr{width:104px;height:104px;flex:none;background:#fff;border-radius:14px;padding:9px}.qr svg{width:100%;height:100%;display:block}
.btns.sm{margin:10px 0 0}.btns.sm .btn{padding:8px 12px;font-size:12px}
.traer .tema{align-items:flex-start;margin-bottom:14px}.traer .tema h3{font:700 19px var(--disp);margin:2px 0 4px}.sub{color:var(--t2);font-size:12.5px;margin:0}
.tags{display:flex;flex-wrap:wrap;gap:7px}.tags span{padding:7px 11px;border-radius:99px;background:var(--fondo);border:1px solid var(--linea);font-size:12.5px;font-weight:600}
.tags.plus{margin-top:8px}.tags.plus span{background:var(--chip);border-color:transparent;color:var(--ink)}
.hint{display:flex;align-items:center;gap:6px;color:var(--t2);font-size:12px;margin:-4px 0 10px}.hint .i{width:14px;height:14px;color:var(--ink)}
a.am{text-decoration:none;color:inherit;cursor:pointer}a.am:active{transform:scale(.97)}
.cat .cab{text-decoration:none;color:inherit}.todo{margin-left:auto;font-size:11.5px;font-weight:600;color:var(--ink)}.todo+.cv{margin-left:2px}
a.lugar{text-decoration:none;color:inherit}a.lugar::after{content:"›";margin-left:auto;padding-left:10px;color:var(--t2)}
.rr .i{color:rgba(255,255,255,.55)!important}.dr{color:inherit}
.cap .cv{color:var(--on)}
/* carrusel */
.car{position:relative;border-radius:18px;overflow:hidden;margin:0 0 14px;background:var(--azul)}
.slides{display:flex;overflow-x:auto;scroll-snap-type:x mandatory;scrollbar-width:none;-webkit-overflow-scrolling:touch}.slides::-webkit-scrollbar{display:none}
.sl{flex:none;width:100%;height:var(--h,210px);scroll-snap-align:center;background:var(--azul) center/cover;position:relative}
.cnt{position:absolute;right:10px;bottom:10px;background:rgba(20,31,58,.65);color:#fff;font:600 11px var(--txt);padding:4px 9px;border-radius:99px}
.sl.ph::before{top:50%}
.credf{font-size:10.5px!important;color:var(--t2);margin:-8px 0 12px!important}.credd{display:block;font-size:9.5px;opacity:.7;margin-top:8px}
.nv{position:absolute;top:50%;transform:translateY(-50%);z-index:2;width:36px;height:36px;border-radius:50%;border:0;background:rgba(255,253,248,.88);color:var(--tinta);display:grid;place-items:center;cursor:pointer;box-shadow:0 2px 8px rgba(20,31,58,.18)}
.nv .i{width:17px;height:17px}.nv.prev{left:10px}.nv.prev .i{transform:rotate(180deg)}.nv.next{right:10px}.nv:active{transform:translateY(-50%) scale(.92)}
.btn.urg{background:#9E1B14;color:#fff;border-color:#9E1B14}
.subt{color:var(--t2);font-size:13px;margin:-4px 0 10px!important}
.pills{display:flex;flex-wrap:wrap;gap:7px;margin:12px 0 4px}
.pill{display:inline-flex;align-items:center;gap:6px;padding:9px 13px;border-radius:99px;background:var(--chip);color:var(--ink);font:600 12.5px var(--txt);text-decoration:none}.pill .i{width:15px;height:15px}
.op{background:var(--tarjeta);border:1px solid var(--linea);border-radius:18px;padding:14px 16px;margin-bottom:10px}
.oph{display:flex;align-items:center;text-decoration:none;color:inherit}.op h4{font:600 15.5px var(--txt);margin:0}.op small{display:block;color:var(--t2);font-size:12.5px;margin-top:3px}.op p{font-size:13.5px;margin:8px 0 0!important}
.cnt2{color:var(--ink);margin:16px 0 10px}
a.pn{text-decoration:none;color:inherit;flex:1}
.filt{display:flex;gap:7px;overflow-x:auto;margin:0 -20px 14px;padding:0 20px 2px;scrollbar-width:none}.filt::-webkit-scrollbar{display:none}
.filt button{flex:none;border:1px solid var(--linea);background:var(--tarjeta);border-radius:99px;padding:8px 13px;font:600 12px var(--txt);cursor:pointer}
.filt button.on{background:var(--tinta);color:var(--crema);border-color:var(--tinta)}
/* galería */
.gsheet .panel{height:88vh;height:88svh}.gsheet .sh-b{display:flex;flex-direction:column;padding-bottom:0}
.gal{flex:1}.gal .sl{height:min(44vh,380px)}.gdesc{font-size:14px;line-height:1.55;margin:0 0 12px}
.zonas-w{position:sticky;bottom:0;background:var(--fondo);border-top:1px solid var(--linea);margin:0 -20px;padding:10px 0 calc(14px + env(safe-area-inset-bottom,0px))}
.zt{padding:0 20px;color:var(--t2);margin:0 0 8px}
.zonas{display:flex;gap:6px;overflow-x:auto;padding:0 20px;scrollbar-width:none;scroll-behavior:smooth}.zonas::-webkit-scrollbar{display:none}
.zn{flex:none;width:74px;display:flex;flex-direction:column;align-items:center;gap:6px;background:none;border:0;padding:0;cursor:pointer;font:600 10.5px/1.2 var(--txt);color:var(--t2);text-align:center}
.zn span{width:50px;height:50px;border-radius:16px;display:grid;place-items:center;background:var(--tarjeta);border:1px solid var(--linea);color:var(--ink)}.zn .i{width:20px;height:20px}
.zn.on{color:var(--tinta)}.zn.on span{background:var(--acc);border-color:var(--acc);color:var(--on)}
.miss{display:none}body.rev .miss{display:inline;background:#FFF4DC;color:#7a4b00;border-radius:4px;padding:0 4px;font-size:.8em}
.revbtn{display:none}body.rev .revbtn{display:flex;position:fixed;top:calc(12px + env(safe-area-inset-top,0px));right:max(12px,calc(50% - 203px));z-index:25;align-items:center;gap:6px;background:#FFF4DC;color:#7a4b00;border:1px dashed #E3B45C;border-radius:99px;padding:7px 12px;font:600 12px var(--txt);text-decoration:none}
"""
JS4 = r"""
const views=[...document.querySelectorAll('[data-view]')];let openS=null,pushed=0;
if(location.search.includes('pendientes'))document.body.classList.add('rev');
const AM=JSON.parse(document.getElementById('amen-data').textContent);
function counter(c){const s=c.querySelector('.slides'),n=c.querySelector('.cnt');if(!n)return;const k=s.children.length;n.textContent=(Math.round(s.scrollLeft/s.clientWidth)+1)+' / '+k;}
const CHV='<svg class="i" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M9 6l6 6-6 6"/></svg>';
const RM=matchMedia('(prefers-reduced-motion: reduce)').matches;
function step(c,d){const s=c.querySelector('.slides'),k=s.children.length;if(k<2)return;let i=Math.round(s.scrollLeft/s.clientWidth)+d;if(i>=k)i=0;if(i<0)i=k-1;s.scrollTo({left:i*s.clientWidth,behavior:RM?'auto':'smooth'});}
function arm(c){const k=c.querySelector('.slides').children.length;
 if(!c.querySelector('.nv')){c.insertAdjacentHTML('beforeend','<button class="nv prev" aria-label="Foto anterior">'+CHV+'</button><button class="nv next" aria-label="Foto siguiente">'+CHV+'</button>');
  c.querySelector('.prev').addEventListener('click',ev=>{ev.stopPropagation();ev.preventDefault();c.dataset.t=Date.now();step(c,-1);});
  c.querySelector('.next').addEventListener('click',ev=>{ev.stopPropagation();ev.preventDefault();c.dataset.t=Date.now();step(c,1);});
  c.querySelector('.slides').addEventListener('pointerdown',()=>c.dataset.t=Date.now(),{passive:true});}
 c.querySelectorAll('.nv').forEach(b=>b.hidden=k<2);}
document.querySelectorAll('.car').forEach(arm);
if(!RM)setInterval(()=>{if(document.hidden)return;document.querySelectorAll('.car').forEach(c=>{if(!c.offsetParent||c.querySelector('.slides').children.length<2)return;if(Date.now()-(+c.dataset.t||0)<9000)return;step(c,1);});},4500);
document.querySelectorAll('.car').forEach(c=>c.querySelector('.slides').addEventListener('scroll',()=>counter(c),{passive:true}));
function gal(i){i=+i||0;const s=document.getElementById('sh-galeria'),z=AM[i];if(!z)return;s.querySelector('.sh-h h2').textContent=z.l;
 const sl=s.querySelector('.slides');sl.innerHTML=z.f.length?z.f.map(u=>'<div class="sl" style="background-image:url('+u+')"></div>').join(''):'<div class="sl ph" data-ph="Foto pendiente · '+z.l.replace(/"/g,'')+'"></div>';sl.scrollLeft=0;
 let c=s.querySelector('.cnt');if(!c){c=document.createElement('span');c.className='cnt';s.querySelector('.car').appendChild(c);}c.hidden=z.f.length<2;c.textContent='1 / '+z.f.length;
 s.querySelector('.gdesc').textContent=z.d;arm(s.querySelector('.car'));s.querySelector('.car').dataset.t=Date.now();
 s.querySelectorAll('.zn').forEach((b,j)=>{b.classList.toggle('on',j===i);if(j===i){const w=b.parentElement;w.scrollLeft=b.offsetLeft-w.clientWidth/2+b.clientWidth/2;}});}
function show(id){const [n,arg]=id.split('.');const s=document.getElementById('sh-'+n);if(!s)return;
 if(openS&&openS!==s)openS.hidden=true;s.hidden=false;openS=s;document.body.classList.add('lock');if(n==='galeria')gal(arg);else s.querySelector('.sh-b').scrollTop=0;}
function hide(){if(openS){openS.hidden=true;openS=null;document.body.classList.remove('lock');}}
function go(){const h=decodeURIComponent(location.hash.slice(1))||'portada';const [vp,sh]=h.split('/');const [v,anc]=vp.split(':');
 const view=views.find(x=>x.dataset.view===v)||views[0];const changed=document.body.dataset.v!==view.dataset.view;
 if(changed){views.forEach(x=>x.hidden=x!==view);document.body.dataset.v=view.dataset.view;}
 document.querySelectorAll('.nav a').forEach(a=>a.classList.toggle('on',a.dataset.v===view.dataset.view));
 if(sh)show(sh);else{hide();const t=anc&&document.getElementById(anc);if(t)setTimeout(()=>t.scrollIntoView({block:'start'}),0);else if(changed)window.scrollTo(0,0);}}
addEventListener('hashchange',go);go();
document.addEventListener('click',ev=>{
 const z=ev.target.closest('.zn');if(z){ev.preventDefault();history.replaceState(null,'','#'+document.body.dataset.v+'/galeria.'+z.dataset.z);gal(z.dataset.z);return;}
 const f=ev.target.closest('.filt button');if(f){const box=f.closest('.sh-b');box.querySelectorAll('.filt button').forEach(b=>b.classList.toggle('on',b===f));
  box.querySelectorAll('.pl[data-cat]').forEach(r=>r.hidden=!(f.dataset.f==='*'||r.dataset.cat===f.dataset.f));return;}
 const t=ev.target.closest('[data-sheet]');
 if(t&&!ev.target.closest('.ab,.pill,.btn:not([data-sheet]),.cp')){ev.preventDefault();pushed++;location.hash=document.body.dataset.v+'/'+t.dataset.sheet;return;}
 if(ev.target.closest('.x')||ev.target.classList.contains('sheet')){ev.preventDefault();location.replace('#'+document.body.dataset.v);}});
addEventListener('keydown',ev=>{if(ev.key==='Escape'&&openS)location.replace('#'+document.body.dataset.v);});
document.querySelectorAll('[data-copy]').forEach(b=>b.addEventListener('click',async ev=>{ev.stopPropagation();try{await navigator.clipboard.writeText(b.dataset.copy);}catch(e){}
 const t=document.querySelector('.toast');t.textContent='Copiado: '+b.dataset.copy;t.classList.add('show');setTimeout(()=>t.classList.remove('show'),1600);}));
if(window.qrcode){if(qrcode.stringToBytesFuncs&&qrcode.stringToBytesFuncs['UTF-8'])qrcode.stringToBytes=qrcode.stringToBytesFuncs['UTF-8'];
 document.querySelectorAll('[data-qr]').forEach(el=>{const q=qrcode(0,'M');q.addData(el.dataset.qr);q.make();el.innerHTML=q.createSvgTag({cellSize:4,margin:0,scalable:true});});}
function toast(m,ms){const t=document.querySelector('.toast');t.textContent=m;t.classList.add('show');setTimeout(()=>t.classList.remove('show'),ms||1800);}
if('serviceWorker' in navigator&&location.protocol.startsWith('http')){navigator.serviceWorker.register('sw.js').catch(()=>{});
 navigator.serviceWorker.addEventListener('message',e=>{if(e.data==='offline-ok')toast('✓ Guía guardada: ya funciona sin conexión',3500);});}
if(location.search.includes('check')){const R={view:document.body.dataset.v,sheet:openS?openS.id:'',over:[],qr:document.querySelectorAll('.qr svg').length,errores:window.__err||[],vh:innerHeight,cta:Math.round(document.querySelector('.abrir').getBoundingClientRect().bottom)};
 const box=(openS?openS.querySelector('.panel'):document.querySelector('.app')).getBoundingClientRect();
 (openS?openS:views.find(x=>!x.hidden)).querySelectorAll('*').forEach(el=>{const r=el.getBoundingClientRect();if(r.width&&(r.right>box.right+1||r.left<box.left-1)&&!el.closest('.slides,.zonas,.filt,.guia')&&getComputedStyle(el).position!=='fixed')R.over.push((el.className||el.tagName)+':'+(el.textContent||'').trim().slice(0,30)+' +'+Math.round(r.right-box.right));});
 R.over=R.over.slice(0,12);const pre=document.createElement('pre');pre.id='check';pre.textContent=JSON.stringify(R);document.body.appendChild(pre);}
"""

SW = r"""const V='vms-__VER__';const PRE=__LIST__;
self.addEventListener('install',e=>{e.waitUntil(caches.open(V).then(c=>c.addAll(PRE)).then(()=>self.skipWaiting()));});
self.addEventListener('activate',e=>{e.waitUntil(caches.keys().then(ks=>Promise.all(ks.filter(k=>k!==V).map(k=>caches.delete(k))))
 .then(()=>self.clients.claim()).then(()=>self.clients.matchAll()).then(cs=>cs.forEach(c=>c.postMessage('offline-ok'))));});
const timeout=(p,ms)=>new Promise((ok,ko)=>{const t=setTimeout(()=>ko('lento'),ms);p.then(r=>{clearTimeout(t);ok(r)},ko);});
self.addEventListener('fetch',e=>{const r=e.request;if(r.method!=='GET')return;const u=new URL(r.url);if(u.origin!==location.origin)return;
 if(r.mode==='navigate'){e.respondWith(timeout(fetch(r),3500).then(res=>{const cp=res.clone();caches.open(V).then(c=>c.put('./',cp));return res;})
  .catch(()=>caches.match('./',{ignoreSearch:true}).then(m=>m||fetch(r))));return;}
 e.respondWith(caches.match(r,{ignoreSearch:true}).then(m=>m||fetch(r).then(res=>{if(res.ok){const cp=res.clone();caches.open(V).then(c=>c.put(r,cp));}return res;})));});
"""
def pwa():
    """Manifiesto + service worker: al abrir la guía una vez con conexión, queda guardada entera y funciona sin señal."""
    import hashlib
    src = os.path.join(W, "_fuente", "pwa")
    os.makedirs(os.path.join(OUT, "assets", "pwa"), exist_ok=True)
    for f in ("icon-192.png", "icon-512.png", "apple-touch-icon.png"):
        if os.path.exists(os.path.join(src, f)): shutil.copy(os.path.join(src, f), os.path.join(OUT, "assets", "pwa", f))
    man = {"name": f"{G.get('nombre', 'Villa Mi Sueño')} · Welcome book", "short_name": G.get("nombre", "VMS"), "start_url": "./#inicio", "scope": "./",
           "display": "standalone", "background_color": "#FBF8F1", "theme_color": "#141F3A", "lang": "es",
           "icons": [{"src": "../assets/pwa/icon-192.png", "sizes": "192x192", "type": "image/png"}, {"src": "../assets/pwa/icon-512.png", "sizes": "512x512", "type": "image/png", "purpose": "any maskable"}]}
    open(os.path.join(OUT, "vms", "manifest.webmanifest"), "w", encoding="utf-8").write(json.dumps(man, ensure_ascii=False))
    pre, h, tot = ["./", "manifest.webmanifest"], hashlib.sha1(), 0
    for root, _d, files in os.walk(OUT):
        if os.sep + "vb" in root: continue
        for f in sorted(files):
            p = os.path.join(root, f); rel = os.path.relpath(p, os.path.join(OUT, "vms")).replace(os.sep, "/")
            if f == "sw.js" or os.path.getsize(p) > 3_000_000 or f.endswith(".pdf") or rel == "../index.html": continue
            h.update(open(p, "rb").read()); tot += os.path.getsize(p)
            if rel not in ("index.html", "manifest.webmanifest"): pre.append(rel)
    sw = SW.replace("__VER__", h.hexdigest()[:10]).replace("__LIST__", json.dumps([urllib.request.quote(x, safe="/.:-_~") for x in pre]))
    open(os.path.join(OUT, "vms", "sw.js"), "w", encoding="utf-8").write(sw)
    print(f"PWA: {len(pre)} archivos para uso sin conexión · aprox. {tot / 1048576:.1f} MB")

def vendor_qr():
    p = os.path.join(VENDOR, "qrcode.min.js")
    if not os.path.exists(p):
        os.makedirs(VENDOR, exist_ok=True)
        open(p, "wb").write(urllib.request.urlopen("https://cdnjs.cloudflare.com/ajax/libs/qrcode-generator/1.4.4/qrcode.min.js").read())
    return open(p, encoding="utf-8").read()

def build():
    load()
    shutil.rmtree(OUT, ignore_errors=True)
    for d in ("assets/fotos", "vms", "vb"): os.makedirs(os.path.join(OUT, d))
    if not os.path.exists(os.path.join(v3.FONTS, "fonts.css")): v3.local_fonts()
    shutil.copytree(v3.FONTS, os.path.join(OUT, "assets", "fonts"), dirs_exist_ok=True)
    fontcss = open(os.path.join(v3.FONTS, "fonts.css")).read()
    logo = open(os.path.join(R, "00_BRAND", "LOGOS_SVG", "vms_logo.svg"), encoding="utf-8").read()
    open(os.path.join(OUT, "assets", "vms_mark.svg"), "w", encoding="utf-8").write('<svg xmlns="http://www.w3.org/2000/svg" viewBox="254 346 87 118">' + "".join(re.findall(r'<path fill-rule="nonzero"[^>]*/>', logo)) + "</svg>")
    pdf = os.path.join(W, "_fuente", "VMS WELCOME PACK-2026_red.pdf")
    copy_svg = "url(\"data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 24 24' fill='none' stroke='black' stroke-width='2'><rect x='8' y='8' width='12' height='12' rx='2'/><path d='M16 8V5a1 1 0 0 0-1-1H5a1 1 0 0 0-1 1v10a1 1 0 0 0 1 1h3'/></svg>\")"
    body = s_portada() + s_inicio() + s_casa() + s_reglas() + s_explora()
    sh = sheets()
    nav = "".join(f'<a href="#{v}" data-v="{v}"><span>{ic(i)}</span>{e(tx)}</a>' for v, tx, i in NAV)
    amen_json = json.dumps(AMEN, ensure_ascii=False).replace("</", "<\\/")
    page = ('<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">'
            f'<meta name="theme-color" content="#141F3A"><meta name="robots" content="noindex"><title>{e(G.get("nombre", ""))} · Welcome book</title><link rel="icon" href="../assets/vms_mark.svg">'
            '<link rel="manifest" href="manifest.webmanifest"><link rel="apple-touch-icon" href="../assets/pwa/apple-touch-icon.png">'
            '<meta name="apple-mobile-web-app-capable" content="yes"><meta name="apple-mobile-web-app-status-bar-style" content="black-translucent"><meta name="apple-mobile-web-app-title" content="Villa Mi Sueño">'
            f'<style>{fontcss}:root{{--copy:{copy_svg}}}{v3.CSS}{CSS4}</style></head><body><div class="app">{body}'
            f'<nav class="nav" aria-label="Secciones">{nav}</nav>{sh}<a class="revbtn" href="#" data-sheet="pendientes">{ic("info")}Pendientes ({len(set(PEND))})</a>'
            f'<div class="toast" role="status"></div></div><script type="application/json" id="amen-data">{amen_json}</script>'
            f'<script>window.__err=[];addEventListener("error",ev=>window.__err.push(ev.message));</script>'
            f'<script>{vendor_qr()}</script><script>{JS4}</script></body></html>')
    open(os.path.join(OUT, "vms", "index.html"), "w", encoding="utf-8").write(page)
    pwa()
    open(os.path.join(OUT, "index.html"), "w", encoding="utf-8").write('<!doctype html><meta http-equiv="refresh" content="0;url=vms/">')
    vb = os.path.join(W, "sitio", "vb", "index.html")
    if os.path.exists(vb): shutil.copy(vb, os.path.join(OUT, "vb", "index.html"))
    print(f"ok -> {os.path.join(OUT, 'vms', 'index.html')}  |  contenido: {XLSX}  |  lugares: {len(D['Lugares'])}  |  pendientes: {len(set(PEND))}")

if __name__ == "__main__":
    build()
